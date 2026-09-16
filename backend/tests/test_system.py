from io import BytesIO
from uuid import uuid4

import pytest
from app import main
from app.extraction import MAX_BYTES, extract
from app.models import EvidenceAnalysis, Report, Run, Source
from app.provider import ProviderError
from app.sample import DemoProvider, sample_case
from app.store import Store
from app.workflow import NODES, audit_references, execute_run
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter


@pytest.fixture
def store(tmp_path):
    return Store(tmp_path)


@pytest.fixture
def client(store, monkeypatch):
    monkeypatch.setattr(main, "store", store)
    monkeypatch.setattr(main.settings, "app_access_token", "")
    with TestClient(main.app) as client:
        yield client


def test_case_validation_and_missing_resources(client):
    assert client.post("/api/cases", json={"title": "x", "scenario": "short"}).status_code == 422
    assert client.get("/api/cases/missing").status_code == 404
    assert client.get("/api/runs/missing").status_code == 404
    case = client.post(
        "/api/cases",
        json={
            "title": "Example contract dispute",
            "scenario": "A fictional supplier delivered goods and claims that an invoice was not paid.",
        },
    ).json()
    assert client.get("/api/cases/" + case["id"]).json()["title"] == case["title"]
    assert client.post(f"/api/cases/{case['id']}/runs", json={"mode": "demo"}).status_code == 422
    assert (
        client.post(f"/api/cases/{case['id']}/runs", json={"mode": "live", "consent": False}).status_code
        == 422
    )


def test_upload_duplicate_and_demo_invalidation(client):
    case = client.post("/api/sample").json()
    url = f"/api/cases/{case['id']}/upload"
    payload = {
        "files": {
            "file": (
                "record.txt",
                b"Fictional invoice with a delivery date and unpaid balance.",
                "text/plain",
            )
        }
    }
    response = client.post(url, **payload)
    assert response.status_code == 201
    assert not response.json()["is_sample"]
    assert len(response.json()["evidence"]) == 4
    assert client.post(url, **payload).status_code == 409
    assert client.post(url, files={"file": ("bad.exe", b"not an executable file")}).status_code == 422
    assert client.post(f"/api/cases/{case['id']}/runs", json={"mode": "demo"}).status_code == 422


def test_upload_unknown_case_and_invalid_kind(client):
    assert (
        client.post(
            "/api/cases/nope/upload", files={"file": ("file.txt", b"some sample textual evidence")}
        ).status_code
        == 404
    )
    case = client.post("/api/sample").json()
    assert (
        client.post(
            f"/api/cases/{case['id']}/upload",
            data={"kind": "invalid"},
            files={"file": ("file.txt", b"some sample textual evidence")},
        ).status_code
        == 422
    )


def test_request_origin_and_token_guard(client, monkeypatch):
    assert client.post("/api/sample", headers={"origin": "https://untrusted.example"}).status_code == 403
    assert client.get("/api/cases", headers={"host": "attacker.example"}).status_code == 400
    monkeypatch.setattr(main.settings, "app_access_token", "test-secret-token")
    assert client.get("/api/cases").status_code == 401
    assert client.get("/api/cases", headers={"Authorization": "Bearer test-secret-token"}).status_code == 200
    assert "test-secret-token" not in client.get("/api/health").text


def test_document_extraction():
    text, note, digest = extract("note.txt", b"Fictional supporting evidence, not a real legal record.")
    assert text.startswith("Fictional") and len(digest) == 64
    doc = Document()
    doc.add_paragraph("Fictional contract with a payment term.")
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Payment within 30 days"
    output = BytesIO()
    doc.save(output)
    assert "Payment within 30 days" in extract("contract.docx", output.getvalue())[0]
    pdf = PdfWriter()
    pdf.add_blank_page(200, 200)
    output = BytesIO()
    pdf.write(output)
    with pytest.raises(ValueError, match="No usable text"):
        extract("scanned.pdf", output.getvalue())
    with pytest.raises(ValueError, match="10 MB"):
        extract("big.txt", b"a" * (MAX_BYTES + 1))
    with pytest.raises(ValueError):
        extract("broken.docx", b"not a valid zip")


async def test_graph_rounds_persistence_and_export(store, client):
    case = sample_case()
    store.save_case(case)
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    await execute_run(run, store)
    saved = store.get_run(run.id)
    assert saved.status == "completed", saved.error
    assert len(saved.rounds) == 2
    assert saved.rounds[0].plaintiff.rebuttals == []
    assert saved.rounds[1].plaintiff.rebuttals
    assert [e.node for e in saved.events if e.status == "completed"] == [key for key, _ in NODES]
    assert saved.report.confidence == "Low"
    assert saved.report.precedents == []
    exported = client.get(f"/api/runs/{run.id}/export")
    assert exported.status_code == 200
    assert "DEMONSTRATION ONLY" in exported.text
    assert "Adversarial round 2" in exported.text
    assert "SHA-256" in exported.text
    assert client.get(f"/api/runs/{run.id}/export?format=json").json()["snapshot"]["id"] == case.id


async def test_agents_independent_in_round_one_and_informed_in_round_two(store):
    class RecordingProvider(DemoProvider):
        calls = []

        async def structured(self, role, task, context, schema):
            if role.endswith("advocate"):
                self.calls.append((role, len(context["previous_rounds"]), set(context)))
            return await super().structured(role, task, context, schema)

    case = sample_case()
    p = RecordingProvider()
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    await execute_run(run, store, p)
    assert run.status == "completed"
    assert [(r, n) for r, n, _ in p.calls] == [
        ("Plaintiff advocate", 0),
        ("Defense advocate", 0),
        ("Plaintiff advocate", 1),
        ("Defense advocate", 1),
    ]
    assert all("current_submissions" not in keys for _, _, keys in p.calls)


async def test_provider_failure_preserves_partial_results(store):
    class FailedProvider(DemoProvider):
        async def research(self, *args):
            raise ProviderError("Test provider unavailable")

    case = sample_case()
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    await execute_run(run, store, FailedProvider())
    saved = store.get_run(run.id)
    assert saved.status == "failed"
    assert saved.structured is not None
    assert saved.events[-1].status == "failed"
    assert saved.error == "Test provider unavailable"


async def test_rejects_missing_evidence_analysis(store):
    class MissingProvider(DemoProvider):
        async def structured(self, role, task, context, schema):
            result = await super().structured(role, task, context, schema)
            if schema is EvidenceAnalysis:
                result.findings = []
            return result

    case = sample_case()
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    await execute_run(run, store, MissingProvider())
    assert run.status == "failed" and "evidence" in run.error.lower()


def test_reference_validation_and_safe_urls():
    warnings = []
    value = {"arguments": [{"source_ids": ["S1", "S99"], "evidence_ids": ["E1", "E99"]}]}
    cleaned = audit_references(value, {"S1"}, {"E1"}, warnings)
    assert cleaned["arguments"][0] == {"source_ids": ["S1"], "evidence_ids": ["E1"]}
    assert len(warnings) == 2
    with pytest.raises(ValueError):
        Source(id="S1", title="Unsafe", url="javascript:alert(1)")


def test_restart_marks_incomplete_runs_interrupted(store):
    case = sample_case()
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    store.save_run(run)
    store.interrupt_stale_runs()
    assert store.get_run(run.id).status == "interrupted"


async def test_claimed_authorities_without_source_are_removed(store):
    from app.models import LegalAuthority

    class BadCitationProvider(DemoProvider):
        async def structured(self, role, task, context, schema):
            result = await super().structured(role, task, context, schema)
            if schema is Report:
                result.precedents = [
                    LegalAuthority(
                        title="Fabricated citation",
                        provision_or_holding="Unsupported",
                        application="Unsupported",
                        source_ids=["S999"],
                        verification_needed="Test",
                    )
                ]
            return result

    case = sample_case()
    run = Run(id=str(uuid4()), case_id=case.id, mode="demo", snapshot=case, model="fixture")
    await execute_run(run, store, BadCitationProvider())
    assert run.report.precedents == []
    assert any("Fabricated citation" in w for w in run.citation_warnings)
