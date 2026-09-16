import hashlib
from io import BytesIO
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

MAX_BYTES = 10 * 1024 * 1024
MAX_CHARS = 60000


def extract(filename: str, content: bytes) -> tuple[str, str, str]:
    if not content or len(content) > MAX_BYTES:
        raise ValueError("Upload a non-empty file smaller than 10 MB.")
    suffix = Path(filename).suffix.lower()
    note = ""
    try:
        if suffix in (".txt", ".md", ".csv"):
            text = content.decode("utf-8-sig")
        elif suffix == ".pdf":
            reader = PdfReader(BytesIO(content))
            if reader.is_encrypted:
                raise ValueError("Password-protected PDFs are unsupported. Upload an unlocked copy.")
            if len(reader.pages) > 150:
                raise ValueError("Upload a PDF with at most 150 pages.")
            pages = [page.extract_text() or "" for page in reader.pages]
            text = "\n\n".join(f"[Page {i + 1}]\n{page}" for i, page in enumerate(pages) if page.strip())
            if any(not page.strip() for page in pages):
                note = "Some PDF pages have no extractable text. Images and scanned pages require OCR before upload."
        elif suffix == ".docx":
            with ZipFile(BytesIO(content)) as archive:
                if sum(item.file_size for item in archive.infolist()) > 40 * 1024 * 1024:
                    raise ValueError("The expanded document exceeds the 40 MB limit.")
            doc = Document(BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs)
            text += "\n" + "\n".join(
                " | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows
            )
            note = "Extracted paragraphs and tables. Embedded images, comments and tracked changes are not assessed."
        else:
            raise ValueError("Supported formats: PDF, DOCX, TXT, MD and CSV.")
    except (ValueError, UnicodeDecodeError, BadZipFile) as exc:
        raise ValueError(str(exc)) from exc
    except Exception as exc:
        raise ValueError("Unable to read this document. Try a text export or an unlocked PDF.") from exc
    text = text.strip().replace("\x00", "")
    if len(text) < 10:
        raise ValueError("No usable text found. For scanned PDFs, run OCR first or paste a transcription.")
    if len(text) > MAX_CHARS:
        raise ValueError("Extracted text exceeds 60,000 characters. Split the document into smaller files.")
    return text, note, hashlib.sha256(content).hexdigest()
